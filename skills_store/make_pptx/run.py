import argparse
import json

from pptx import Presentation
from pptx.util import Pt
from pptx.oxml.ns import qn


def set_cjk(run, size, font, bold=False):
    """Set font size/bold and apply the typeface to latin, East Asian and
    complex-script runs, so Chinese glyphs do not fall back unpredictably."""
    run.font.size = Pt(size)
    run.font.bold = bold
    rPr = run.font._rPr
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", font)


def add_slide(prs, slot, title, bullets, font, title_size, body_size):
    slide = prs.slides.add_slide(prs.slide_layouts[slot])
    slide.shapes.title.text = title
    for run in slide.shapes.title.text_frame.paragraphs[0].runs:
        set_cjk(run, title_size, font, bold=True)
    body = slide.placeholders[1].text_frame
    for i, bullet in enumerate(bullets):
        para = body.paragraphs[0] if i == 0 else body.add_paragraph()
        para.text = bullet
        for run in para.runs:
            set_cjk(run, body_size, font)
    return slide


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--args-json", required=True)
    args = json.loads(parser.parse_args().args_json)

    font = args.get("font", "微软雅黑")
    prs = Presentation()

    # Title slide (layout 0), then content slides (layout 1).
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = args["title"]
    for run in title_slide.shapes.title.text_frame.paragraphs[0].runs:
        set_cjk(run, 40, font, bold=True)

    for item in args["slides"]:
        add_slide(prs, 1, item["title"], item["bullets"], font, 30, 20)

    prs.save(args["output_path"])
    print(f"已生成 {args['output_path']}（共 {len(prs.slides._sldIdLst)} 页）")


if __name__ == "__main__":
    main()
