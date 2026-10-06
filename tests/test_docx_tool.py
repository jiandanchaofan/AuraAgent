"""Tests for tools/documents/docx_tool.py -- read_docx/create_docx/
edit_docx. Mostly create-then-read round trips (no external fixture
files needed) plus the sandbox/size/mode-guard cases every tools/documents/
module shares.
"""
from __future__ import annotations

import docx
import pytest

from core.exceptions import ToolExecutionError
from tools.documents.docx_tool import register_docx_tools
from tools.registry import ToolRegistry
from tools.workspace_root import SwappableWorkspaceRoot


def _registry(tmp_path, max_read_bytes=20_000_000, max_extract_chars=50_000):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    registry = ToolRegistry()
    register_docx_tools(registry, SwappableWorkspaceRoot(workspace), max_read_bytes, max_extract_chars)
    return registry, workspace


@pytest.mark.asyncio
async def test_create_then_read_round_trip(tmp_path):
    registry, _ = _registry(tmp_path)

    await registry.dispatch(
        "create_docx",
        {"path": "report.docx", "title": "Test Report", "paragraphs": ["First paragraph.", "Second paragraph."]},
    )
    result = await registry.dispatch("read_docx", {"path": "report.docx"})

    assert "Test Report" in result
    assert "First paragraph." in result
    assert "Second paragraph." in result


@pytest.mark.asyncio
async def test_read_docx_includes_tables(tmp_path):
    registry, workspace = _registry(tmp_path)
    document = docx.Document()
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Score"
    table.cell(1, 0).text = "Alice"
    table.cell(1, 1).text = "9"
    document.save(str(workspace / "with_table.docx"))

    result = await registry.dispatch("read_docx", {"path": "with_table.docx"})

    assert "--- Tables ---" in result
    assert "Name | Score" in result
    assert "Alice | 9" in result


@pytest.mark.asyncio
async def test_create_docx_saves_within_workspace_with_subdirectory(tmp_path):
    registry, workspace = _registry(tmp_path)

    result = await registry.dispatch("create_docx", {"path": "reports/deck.docx", "paragraphs": ["hi"]})

    saved_path = workspace / "reports" / "deck.docx"
    assert saved_path.is_file()
    assert "reports/deck.docx" in result


@pytest.mark.asyncio
async def test_create_docx_create_only_fails_if_exists(tmp_path):
    registry, _ = _registry(tmp_path)
    await registry.dispatch("create_docx", {"path": "report.docx", "paragraphs": ["a"]})

    with pytest.raises(ToolExecutionError, match="already exists"):
        await registry.dispatch("create_docx", {"path": "report.docx", "paragraphs": ["b"]})


@pytest.mark.asyncio
async def test_create_docx_overwrite_replaces(tmp_path):
    registry, _ = _registry(tmp_path)
    await registry.dispatch("create_docx", {"path": "report.docx", "paragraphs": ["original"]})

    await registry.dispatch("create_docx", {"path": "report.docx", "paragraphs": ["replaced"], "mode": "overwrite"})
    result = await registry.dispatch("read_docx", {"path": "report.docx"})

    assert "replaced" in result
    assert "original" not in result


@pytest.mark.asyncio
async def test_edit_docx_appends_paragraphs(tmp_path):
    registry, _ = _registry(tmp_path)
    await registry.dispatch("create_docx", {"path": "report.docx", "paragraphs": ["original"]})

    await registry.dispatch("edit_docx", {"path": "report.docx", "append_paragraphs": ["appended one", "appended two"]})
    result = await registry.dispatch("read_docx", {"path": "report.docx"})

    assert "original" in result
    assert "appended one" in result
    assert "appended two" in result


@pytest.mark.asyncio
async def test_edit_docx_find_replace(tmp_path):
    registry, _ = _registry(tmp_path)
    await registry.dispatch("create_docx", {"path": "report.docx", "paragraphs": ["The quick brown fox."]})

    await registry.dispatch("edit_docx", {"path": "report.docx", "find": "quick brown fox", "replace": "lazy dog"})
    result = await registry.dispatch("read_docx", {"path": "report.docx"})

    assert "The lazy dog." in result
    assert "quick brown fox" not in result


@pytest.mark.asyncio
async def test_edit_docx_no_changes_when_nothing_to_do(tmp_path):
    registry, _ = _registry(tmp_path)
    await registry.dispatch("create_docx", {"path": "report.docx", "paragraphs": ["original"]})

    result = await registry.dispatch("edit_docx", {"path": "report.docx", "find": "not present"})

    assert "No changes made" in result


@pytest.mark.asyncio
async def test_edit_docx_missing_file(tmp_path):
    registry, _ = _registry(tmp_path)

    with pytest.raises(ToolExecutionError, match="File not found"):
        await registry.dispatch("edit_docx", {"path": "missing.docx", "append_paragraphs": ["x"]})


@pytest.mark.asyncio
async def test_read_docx_rejects_oversized_file(tmp_path):
    registry, _ = _registry(tmp_path, max_read_bytes=10)
    await registry.dispatch("create_docx", {"path": "report.docx", "paragraphs": ["hello world, this is long enough"]})

    with pytest.raises(ToolExecutionError, match="over the"):
        await registry.dispatch("read_docx", {"path": "report.docx"})


@pytest.mark.asyncio
async def test_read_docx_truncates_long_output(tmp_path):
    registry, _ = _registry(tmp_path, max_extract_chars=10)
    await registry.dispatch("create_docx", {"path": "report.docx", "paragraphs": ["a much longer paragraph than the cap"]})

    result = await registry.dispatch("read_docx", {"path": "report.docx"})

    assert "truncated" in result


@pytest.mark.asyncio
async def test_read_docx_blocks_path_traversal(tmp_path):
    registry, _ = _registry(tmp_path)

    with pytest.raises(ToolExecutionError, match="outside the sandbox root"):
        await registry.dispatch("read_docx", {"path": "../escaped.docx"})


@pytest.mark.asyncio
async def test_create_docx_rejects_absolute_path(tmp_path):
    registry, _ = _registry(tmp_path)
    absolute_target = tmp_path / "elsewhere.docx"

    with pytest.raises(ToolExecutionError, match="Absolute paths are not allowed"):
        await registry.dispatch("create_docx", {"path": str(absolute_target), "paragraphs": ["x"]})
    assert not absolute_target.exists()
