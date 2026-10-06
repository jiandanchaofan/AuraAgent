"""read_pptx/create_pptx/edit_pptx -- PowerPoint deck read/create/edit via
python-pptx. create_pptx supersedes the earlier skills_store/make_pptx
Skill -- _set_cjk()/_add_slide() below are migrated verbatim from that
Skill's run.py (the CJK-typeface handling in particular was a real,
hard-won fix; not worth re-deriving). As a native tool this reuses
resolve_within_sandbox()/workspace_root directly instead of the Skill's
own AURA_WORKSPACE_ROOT-based workaround (skills/skill_loader.py runs a
Skill's subprocess with the Skill's OWN directory as cwd, not the
workspace -- a native tool never has that problem, since it runs
in-process).
"""
from __future__ import annotations

from typing import Any

from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Pt

from core.exceptions import ToolExecutionError
from tools.base import ToolSpec
from tools.documents._common import check_readable_size, truncate_with_note
from tools.registry import ToolRegistry
from tools.sandbox_path import resolve_within_sandbox
from tools.workspace_root import SwappableWorkspaceRoot

_DEFAULT_FONT = "微软雅黑"


def _set_cjk(run, size, font, bold=False) -> None:
    """Set font size/bold and apply the typeface to latin, East Asian and
    complex-script runs, so Chinese glyphs do not fall back unpredictably."""
    run.font.size = Pt(size)
    run.font.bold = bold
    rPr = run.font._rPr
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", font)


def _add_slide(prs, slot, title, bullets, font, title_size, body_size):
    slide = prs.slides.add_slide(prs.slide_layouts[slot])
    slide.shapes.title.text = title
    for run in slide.shapes.title.text_frame.paragraphs[0].runs:
        _set_cjk(run, title_size, font, bold=True)
    body = slide.placeholders[1].text_frame
    for i, bullet in enumerate(bullets):
        para = body.paragraphs[0] if i == 0 else body.add_paragraph()
        para.text = bullet
        for run in para.runs:
            _set_cjk(run, body_size, font)
    return slide


def _read_pptx_text(path) -> str:
    prs = Presentation(str(path))
    parts = []
    for i, slide in enumerate(prs.slides, start=1):
        title = slide.shapes.title.text if slide.shapes.title else ""
        lines = []
        for shape in slide.shapes:
            if shape == slide.shapes.title or not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs)
                if text:
                    lines.append(f"- {text}")
        parts.append(f"--- Slide {i}: {title} ---\n" + "\n".join(lines))
    return "\n\n".join(parts)


def register_pptx_tools(
    registry: ToolRegistry,
    workspace_root: SwappableWorkspaceRoot,
    max_read_bytes: int,
    max_extract_chars: int,
) -> None:
    async def read_pptx(args: dict[str, Any]) -> str:
        path = resolve_within_sandbox(workspace_root.current, args["path"])
        if not path.is_file():
            raise ToolExecutionError(f"File not found: '{args['path']}'")
        check_readable_size(path, max_read_bytes)
        try:
            text = _read_pptx_text(path)
        except Exception as exc:  # noqa: BLE001 - python-pptx raises several different exception types for malformed input
            raise ToolExecutionError(f"Could not open '{args['path']}' as a .pptx: {exc}") from exc
        return truncate_with_note(text, max_extract_chars)

    registry.register(
        ToolSpec(
            name="read_pptx",
            description="Read the text content of a .pptx (PowerPoint) file in the workspace -- each slide's title and bullet text.",
            input_schema={
                "type": "object",
                "properties": {"path": {"type": "string", "description": "File path relative to the workspace."}},
                "required": ["path"],
            },
        ),
        read_pptx,
    )

    async def create_pptx(args: dict[str, Any]) -> str:
        path = resolve_within_sandbox(workspace_root.current, args["path"])
        mode = args.get("mode", "create_only")
        if mode not in ("create_only", "overwrite"):
            raise ToolExecutionError(f"Unknown mode '{mode}' (expected 'create_only' or 'overwrite').")
        if mode == "create_only" and path.exists():
            raise ToolExecutionError(f"'{args['path']}' already exists (use mode='overwrite' to replace it).")

        font = args.get("font", _DEFAULT_FONT)
        prs = Presentation()

        title_slide = prs.slides.add_slide(prs.slide_layouts[0])
        title_slide.shapes.title.text = args["title"]
        for run in title_slide.shapes.title.text_frame.paragraphs[0].runs:
            _set_cjk(run, 40, font, bold=True)

        for item in args.get("slides", []):
            _add_slide(prs, 1, item["title"], item["bullets"], font, 30, 20)

        path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(str(path))
        return f"Created '{args['path']}' ({len(prs.slides._sldIdLst)} slide(s))."

    registry.register(
        ToolSpec(
            name="create_pptx",
            description=(
                "Create a Chinese-friendly PowerPoint (.pptx) deck from a title and a list of content "
                "slides (each with its own title and bullet points). mode='create_only' (default) fails "
                "if the file already exists; mode='overwrite' replaces it."
            ),
            input_schema={
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "File path relative to the workspace."},
                    "title": {"type": "string", "description": "Deck title, rendered on the first slide."},
                    "slides": {
                        "type": "array",
                        "description": "Content slides, in order.",
                        "items": {
                            "type": "object",
                            "properties": {
                                "title": {"type": "string"},
                                "bullets": {"type": "array", "items": {"type": "string"}},
                            },
                            "required": ["title", "bullets"],
                        },
                    },
                    "font": {"type": "string", "description": "CJK typeface name, default '微软雅黑' (Microsoft YaHei)."},
                    "mode": {"type": "string", "enum": ["create_only", "overwrite"], "description": "Defaults to 'create_only'."},
                },
                "required": ["path", "title", "slides"],
            },
        ),
        create_pptx,
    )

    async def edit_pptx(args: dict[str, Any]) -> str:
        path = resolve_within_sandbox(workspace_root.current, args["path"])
        if not path.is_file():
            raise ToolExecutionError(f"File not found: '{args['path']}'")

        append_slides = args.get("append_slides") or []
        if not append_slides:
            return "No changes made (append_slides was empty)."

        prs = Presentation(str(path))
        font = args.get("font", _DEFAULT_FONT)
        for item in append_slides:
            _add_slide(prs, 1, item["title"], item["bullets"], font, 30, 20)

        prs.save(str(path))
        return f"Appended {len(append_slides)} slide(s) to '{args['path']}'."

    registry.register(
        ToolSpec(
            name="edit_pptx",
            description=(
                "Append new content slides to the end of an EXISTING .pptx deck. Not a general editor -- "
                "cannot modify or reorder existing slides, only add new ones after them."
            ),
            input_schema={
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "File path relative to the workspace."},
                    "append_slides": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "title": {"type": "string"},
                                "bullets": {"type": "array", "items": {"type": "string"}},
                            },
                            "required": ["title", "bullets"],
                        },
                    },
                    "font": {"type": "string", "description": "CJK typeface name, default '微软雅黑' (Microsoft YaHei)."},
                },
                "required": ["path", "append_slides"],
            },
        ),
        edit_pptx,
    )
