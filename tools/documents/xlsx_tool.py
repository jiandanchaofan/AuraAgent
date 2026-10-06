"""read_xlsx/create_xlsx/edit_xlsx -- Excel workbook read/create/edit via
openpyxl (pure Python, no external binary dependency -- unlike a
pandas-based approach, this needs nothing beyond openpyxl itself).
"""
from __future__ import annotations

from typing import Any

import openpyxl

from core.exceptions import ToolExecutionError
from tools.base import ToolSpec
from tools.documents._common import check_readable_size, truncate_with_note
from tools.registry import ToolRegistry
from tools.sandbox_path import resolve_within_sandbox
from tools.workspace_root import SwappableWorkspaceRoot


def register_xlsx_tools(
    registry: ToolRegistry,
    workspace_root: SwappableWorkspaceRoot,
    max_read_bytes: int,
    max_extract_chars: int,
) -> None:
    async def read_xlsx(args: dict[str, Any]) -> str:
        path = resolve_within_sandbox(workspace_root.current, args["path"])
        if not path.is_file():
            raise ToolExecutionError(f"File not found: '{args['path']}'")
        check_readable_size(path, max_read_bytes)
        try:
            # data_only=True reads each formula cell's last-computed value
            # (what a human sees in Excel), not the formula text itself.
            workbook = openpyxl.load_workbook(str(path), data_only=True)
        except Exception as exc:  # noqa: BLE001 - openpyxl raises several different exception types for malformed input
            raise ToolExecutionError(f"Could not open '{args['path']}' as a .xlsx: {exc}") from exc

        sheet_name = args.get("sheet")
        if not sheet_name:
            return "Sheets: " + ", ".join(workbook.sheetnames)
        if sheet_name not in workbook.sheetnames:
            raise ToolExecutionError(f"No sheet named '{sheet_name}'. Available: {', '.join(workbook.sheetnames)}")

        worksheet = workbook[sheet_name]
        lines = [
            ", ".join("" if cell is None else str(cell) for cell in row) for row in worksheet.iter_rows(values_only=True)
        ]
        return truncate_with_note("\n".join(lines), max_extract_chars)

    registry.register(
        ToolSpec(
            name="read_xlsx",
            description=(
                "Read a .xlsx (Excel) file in the workspace. Omit `sheet` to list the workbook's sheet "
                "names; pass one to read its rows as comma-separated text (one line per row). Formula "
                "cells return their last-computed value, not the formula itself."
            ),
            input_schema={
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "File path relative to the workspace."},
                    "sheet": {"type": "string", "description": "Sheet name to read. Omit to list all sheet names."},
                },
                "required": ["path"],
            },
        ),
        read_xlsx,
    )

    async def create_xlsx(args: dict[str, Any]) -> str:
        path = resolve_within_sandbox(workspace_root.current, args["path"])
        mode = args.get("mode", "create_only")
        if mode not in ("create_only", "overwrite"):
            raise ToolExecutionError(f"Unknown mode '{mode}' (expected 'create_only' or 'overwrite').")
        if mode == "create_only" and path.exists():
            raise ToolExecutionError(f"'{args['path']}' already exists (use mode='overwrite' to replace it).")

        sheets = args.get("sheets", [])
        if not sheets:
            raise ToolExecutionError("`sheets` must contain at least one sheet.")

        workbook = openpyxl.Workbook()
        workbook.remove(workbook.active)  # drop the default empty "Sheet"
        for sheet in sheets:
            worksheet = workbook.create_sheet(sheet["name"])
            for row in sheet.get("rows", []):
                worksheet.append(row)

        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(str(path))
        return f"Created '{args['path']}' with {len(sheets)} sheet(s)."

    registry.register(
        ToolSpec(
            name="create_xlsx",
            description=(
                "Create a new .xlsx (Excel) file in the workspace from one or more sheets, each with a "
                "name and a list of rows (each row a list of cell values). mode='create_only' (default) "
                "fails if the file already exists; mode='overwrite' replaces it."
            ),
            input_schema={
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "File path relative to the workspace."},
                    "sheets": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "name": {"type": "string"},
                                "rows": {
                                    "type": "array",
                                    "items": {"type": "array", "items": {}},
                                    "description": "Each item is one row: a list of cell values (string/number).",
                                },
                            },
                            "required": ["name"],
                        },
                    },
                    "mode": {"type": "string", "enum": ["create_only", "overwrite"], "description": "Defaults to 'create_only'."},
                },
                "required": ["path", "sheets"],
            },
        ),
        create_xlsx,
    )

    async def edit_xlsx(args: dict[str, Any]) -> str:
        path = resolve_within_sandbox(workspace_root.current, args["path"])
        if not path.is_file():
            raise ToolExecutionError(f"File not found: '{args['path']}'")

        try:
            workbook = openpyxl.load_workbook(str(path))
        except Exception as exc:  # noqa: BLE001 - openpyxl raises several different exception types for malformed input
            raise ToolExecutionError(f"Could not open '{args['path']}' as a .xlsx: {exc}") from exc

        sheet_name = args["sheet"]
        if sheet_name not in workbook.sheetnames:
            raise ToolExecutionError(f"No sheet named '{sheet_name}'. Available: {', '.join(workbook.sheetnames)}")
        worksheet = workbook[sheet_name]

        append_rows = args.get("append_rows") or []
        for row in append_rows:
            worksheet.append(row)

        set_cells = args.get("set_cells") or []
        for entry in set_cells:
            worksheet[entry["cell"]] = entry["value"]

        if not append_rows and not set_cells:
            return "No changes made (nothing to append, no cells to set)."

        workbook.save(str(path))
        return f"Updated sheet '{sheet_name}' in '{args['path']}': appended {len(append_rows)} row(s), set {len(set_cells)} cell(s)."

    registry.register(
        ToolSpec(
            name="edit_xlsx",
            description=(
                "Make a small, well-defined edit to a sheet in an EXISTING .xlsx file: append new rows "
                "to the end, and/or set specific cells by reference (e.g. 'B3')."
            ),
            input_schema={
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "File path relative to the workspace."},
                    "sheet": {"type": "string", "description": "Sheet name to edit."},
                    "append_rows": {
                        "type": "array",
                        "items": {"type": "array", "items": {}},
                        "description": "Each item is one row (a list of cell values) appended at the end.",
                    },
                    "set_cells": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {"cell": {"type": "string"}, "value": {}},
                            "required": ["cell", "value"],
                        },
                        "description": "Individual cells to set, e.g. [{\"cell\": \"B3\", \"value\": 42}].",
                    },
                },
                "required": ["path", "sheet"],
            },
        ),
        edit_xlsx,
    )
